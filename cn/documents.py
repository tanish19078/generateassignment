"""Sample-style CN report, with no cover page or fabricated observations."""
import io
from docx import Document
from docx.shared import Pt,Cm,RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH,WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from .experiments import build_experiment
from .rendering import render_figure


def paragraph(doc,text='',bold=False,size=12,spacing=1.5):
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.line_spacing=spacing;p.paragraph_format.space_after=Pt(5)
    r=p.add_run(text);r.font.name='Times New Roman';r.font.size=Pt(size);r.bold=bold
    return p


def heading(doc,text,major=False):
    p=paragraph(doc,text,True,16 if major else 14,1)
    p.alignment=WD_ALIGN_PARAGRAPH.LEFT;p.paragraph_format.keep_with_next=True
    p.paragraph_format.space_before=Pt(5)
    return p


def table(doc,headers,rows,caption):
    p=paragraph(doc,caption,spacing=1);p.alignment=WD_ALIGN_PARAGRAPH.CENTER;p.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
    widths = ([1.85,2.8,3.7,3.9,3.67] if headers == ['Device','Interface','IP Address','Subnet Mask','Gateway'] else [15.92/len(headers)]*len(headers))
    for column,width in zip(t.columns,widths):column.width=Cm(width)
    for c,label in zip(t.rows[0].cells,headers):c.text=str(label)
    for row in rows:
        for c,value in zip(t.add_row().cells,row):c.text=str(value)
    for i,row in enumerate(t.rows):
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        for j,c in enumerate(row.cells):
            c.width=Cm(widths[j]);pr=c._tc.get_or_add_tcPr();border=OxmlElement('w:tcBorders')
            for edge in ('top','left','bottom','right'):
                b=OxmlElement('w:'+edge);b.set(qn('w:val'),'single');b.set(qn('w:sz'),'4');b.set(qn('w:color'),'000000');border.append(b)
            pr.append(border)
            for p in c.paragraphs:
                p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.line_spacing=1;p.paragraph_format.space_after=Pt(3)
                p.paragraph_format.keep_with_next=i < len(t.rows)-1
                for run in p.runs:run.font.name='Times New Roman';run.font.size=Pt(12);run.bold=i==0
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))


def add_figure(doc,exp,figure,n,index):
    raw=render_figure(exp,figure)
    p=doc.add_paragraph();p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_together=True;p.paragraph_format.space_after=Pt(5)
    with io.BytesIO(raw) as buf:p.add_run().add_picture(buf,width=Cm(15.92))
    p.add_run().add_break()
    r=p.add_run(f"Fig {n}.{index}: {figure['caption']}")
    r.font.name='Times New Roman';r.font.size=Pt(12)


def add_experiment(doc,exp,n,start_new_page=True):
    title = heading(doc,f'Experiment No. {n}',True)
    title.paragraph_format.page_break_before = n > 1 and start_new_page
    details=' | '.join(exp['profile'][key] for key in ('name','roll','section') if exp['profile'][key])
    if details:paragraph(doc,details,spacing=1)
    p=paragraph(doc,spacing=1);r=p.add_run('Aim: ');r.bold=True;r.font.name='Times New Roman';r.font.size=Pt(14)
    r=p.add_run(exp['aim']);r.font.name='Times New Roman';r.font.size=Pt(12)
    heading(doc,'Description');paragraph(doc,exp['concept'])
    heading(doc,'Requirements')
    for item in exp['requirements']:paragraph(doc,'• '+item,spacing=1)
    heading(doc,'Network Topology')
    figure_number=1
    for figure in exp['figures']:
        if figure['kind'] in ('topology','subnets'):
            add_figure(doc,exp,figure,n,figure_number);figure_number+=1
    table_number=0
    if exp['address_rows']:
        table_number+=1
        table(doc,['Device','Interface','IP Address','Subnet Mask','Gateway'],exp['address_rows'],f'Table {n}.{table_number}: IP Address Configuration')
    if exp['topology']['links']:
        table_number+=1
        rows=[[f"{link['a']} {link['aport']}",f"{link['b']} {link['bport']}",link['cable']] for link in exp['topology']['links']]
        table(doc,['From','To','Cable'],rows,f'Table {n}.{table_number}: Device Connections')
    heading(doc,'Procedure/Commands')
    for i,step in enumerate(exp['procedure'],1):paragraph(doc,f'{i}. {step}',spacing=1)
    for cfg in exp['configurations']:
        p=paragraph(doc,cfg['device'],True,spacing=1);p.paragraph_format.keep_with_next=True
        p=paragraph(doc,'\n'.join(cfg['commands']),spacing=1)
        # Long variants may flow over a page; individual lines stay intact.
        p.alignment=WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.left_indent=Cm(.65)
    for data in exp['tables']:
        table_number+=1;table(doc,data['headers'],data['rows'],f"Table {n}.{table_number}: {data['title']}")
    for figure in exp['figures']:
        if figure['kind']=='ip':
            add_figure(doc,exp,figure,n,figure_number);figure_number+=1
    p=paragraph(doc,'Verification commands',True,spacing=1);p.paragraph_format.keep_with_next=True
    paragraph(doc,'\n'.join(exp['verification']),spacing=1).alignment=WD_ALIGN_PARAGRAPH.LEFT
    for figure in exp['figures']:
        if figure['kind'] in ('connectivity','verification'):
            add_figure(doc,exp,figure,n,figure_number);figure_number+=1
    heading(doc,'Result');paragraph(doc,exp['result'])


def configure_document(doc):
    sec=doc.sections[0];sec.page_width=Cm(21);sec.page_height=Cm(29.7)
    sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Cm(2.54)
    st=doc.styles['Normal'];st.font.name='Times New Roman';st.font.size=Pt(12);st.font.color.rgb=RGBColor(0,0,0)
    f=sec.footer.paragraphs[0];f.paragraph_format.tab_stops.clear_all()
    # Override the inherited centre tab stop so page numbers stay at the right.
    f.style=doc.styles['Normal'];f.paragraph_format.tab_stops.add_tab_stop(Cm(15.92),WD_TAB_ALIGNMENT.RIGHT)
    pr=f._p.get_or_add_pPr();border=OxmlElement('w:pBdr');line=OxmlElement('w:top')
    for key,value in [('val','double'),('sz','6'),('space','6'),('color','773535')]:line.set(qn('w:'+key),value)
    border.append(line);pr.append(border)
    r=f.add_run('Computer Networks\tPage ');r.font.name='Times New Roman';r.font.size=Pt(10)
    field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');f._p.append(field)


def canonical_experiment(raw,profile=None):
    """Never accept client-supplied figures as proof of an actual simulator run."""
    return build_experiment(raw.get('aim',''),raw.get('profile') or profile)


def build_cn_document(experiments,settings=None):
    settings=settings or {};doc=Document();configure_document(doc)
    for i,raw in enumerate(experiments,1):
        add_experiment(doc,canonical_experiment(raw,settings.get('student')),i)
    buf=io.BytesIO();doc.save(buf);buf.seek(0);return buf
