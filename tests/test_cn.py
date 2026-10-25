import copy
import ipaddress
import pytest
from cn.experiments import CATALOG, build_experiment, identify_experiment, validate_plan

@pytest.mark.parametrize('key,title,aim', CATALOG)
def test_all_recipes_have_consistent_addresses(key,title,aim):
    exp=build_experiment(aim,{'name':'Tanish','roll':'2510990000'})
    assert exp['template_id']==key
    assert exp['execution']['status']=='not_run'
    assert exp['procedure'] and exp['figures']
    for dev in exp['topology']['devices']:
        for port,value in dev['ports'].items():
            assert ipaddress.ip_interface(value['ip']+'/'+str(value['prefix']))
            assert any(row[0]==dev['id'] and row[2]==value['ip'] for row in exp['address_rows'])
            if dev['kind']=='router':
                config=next(c for c in exp['configurations'] if c['device']==dev['id'])
                assert f"ip address {value['ip']} {value['mask']}" in config['commands']

def test_identity_is_stable_and_personalized():
    a=build_experiment(CATALOG[2][2],{'name':'A','roll':'1'})
    assert a==build_experiment(CATALOG[2][2],{'name':'A','roll':'1'})
    assert a['address_rows']!=build_experiment(CATALOG[2][2],{'name':'B','roll':'2'})['address_rows']

def test_overlapping_protocols_are_not_silently_dropped():
    with pytest.raises(ValueError,match='Separate'):
        build_experiment('Configure RIP and OSPF')
    with pytest.raises(ValueError,match='currently supports'):
        build_experiment('Configure BGP peering')
    assert identify_experiment('Write a Python program for subnetting') is None


def test_link_conflicts_are_rejected():
    plan=build_experiment(CATALOG[1][2])['topology']
    broken=copy.deepcopy(plan);broken['links']*=2
    with pytest.raises(ValueError,match='reused'):
        validate_plan(broken)
    broken=copy.deepcopy(plan);broken['devices'][1]['ports']['FastEthernet0']['ip']='192.168.255.12'
    with pytest.raises(ValueError,match='share a subnet'):
        validate_plan(broken)


def test_vlsm_allocations_do_not_overlap():
    exp=build_experiment(CATALOG[4][2])
    for table in exp['tables'][3:]:
        nets=[ipaddress.ip_network(row[1]) for row in table['rows']]
        for i,left in enumerate(nets):
            for right in nets[i+1:]:assert not left.overlaps(right)
        for row,net in zip(table['rows'],nets):assert net.num_addresses-2>=int(row[0])


def test_figures_decode_as_png_and_do_not_claim_execution():
    import io
    from PIL import Image
    from cn.rendering import render_figure
    exp=build_experiment(CATALOG[2][2],{'name':'Tanish','roll':'42'})
    for figure in exp['figures']:
        with Image.open(io.BytesIO(render_figure(exp,figure))) as image:
            assert image.format=='PNG'
            assert image.width>=1000
    assert 'Expected' in exp['figures'][-1]['caption']


def test_document_rebuilds_canonical_figures_without_cover():
    from cn.documents import build_cn_document
    from docx import Document
    raw=build_experiment(CATALOG[2][2],{'name':'Tanish','roll':'42'})
    raw['execution']={'status':'verified'}
    raw['figures']=[{'kind':'fake','data_url':'file:///etc/passwd'}]
    doc=Document(build_cn_document([raw]))
    text='\n'.join(p.text for p in doc.paragraphs)
    assert doc.paragraphs[0].text=='Experiment No. 1'
    assert 'Tanish | 42' in text and 'Expected' in text
    assert len(doc.inline_shapes)==3
    assert 'Requirements' in text and 'Procedure/Commands' in text


def test_cn_generate_needs_no_provider_or_key(monkeypatch):
    import app
    monkeypatch.setattr(app,'run_completion',lambda *args,**kwargs:pytest.fail('CN must not call an LLM'))
    with app.app.test_client() as client:
        response=client.post('/api/generate',json={'mode':'cn','aim':CATALOG[2][2],'student':{'name':'Tanish'}})
        assert response.status_code==200
        body=response.get_json()
        assert body['mode']=='cn' and body['execution']['status']=='not_run'
        assert all(f['data_url'].startswith('data:image/png;base64,') for f in body['figures'])
        assert client.post('/api/generate',json={'mode':'cn','aim':'Configure BGP'}).status_code==400
        assert client.post('/api/generate',json={'mode':'cn','aim':CATALOG[0][2],'student':{'name':'x\nforged'}}).status_code==400


def test_auto_detect_and_mixed_export():
    import app
    from docx import Document
    assert app.detect_mode_and_language(CATALOG[6][2])['mode']=='cn'
    assert app.detect_mode_and_language('Write a Python program for RIP')['mode']=='general'
    doc=Document(app.build_document([
        {'mode':'cn','aim':CATALOG[1][2],'profile':{'name':'A'}},
        {'aim':'Print hello','concept':'Basic output','code':'print("hello")','output':'hello'}],{},'auto'))
    text='\n'.join(p.text for p in doc.paragraphs)
    assert 'Experiment No. 1' in text and 'Experiment No. 2' in text


def test_plan_export_and_document_endpoint():
    import app
    from docx import Document
    import io
    with app.app.test_client() as client:
        assert len(client.get('/api/cn/catalog').get_json()['experiments'])==8
        plan=client.post('/api/cn/plan',json={'aim':CATALOG[7][2]}).get_json()
        assert 'data_url' not in plan['figures'][0]
        assert plan['topology']['devices']
        response=client.post('/api/download',json={'mode':'cn','experiments':[{'aim':CATALOG[2][2],'profile':{'name':'A'}}]})
        assert response.status_code==200
        assert len(Document(io.BytesIO(response.data)).inline_shapes)==3


def test_cn_export_budget_uses_actual_figure_count_and_rejects_unknown_aims():
    import app
    with app.app.test_client() as client:
        response=client.post('/api/download',json={'mode':'cn','experiments':[{'aim':CATALOG[7][2]}]*51})
        assert response.status_code==400 and 'Too many' in response.get_json()['error']
        response=client.post('/api/download',json={'mode':'cn','experiments':[{'aim':'Unsupported BGP lab'}]})
        assert response.status_code==400


def test_command_blocks_are_left_aligned_and_experiments_use_heading_breaks():
    from cn.documents import build_cn_document
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    doc=Document(build_cn_document([{'aim':CATALOG[6][2]},{'aim':CATALOG[7][2]}]))
    commands=[p for p in doc.paragraphs if 'configure terminal\n' in p.text]
    assert commands and all(p.alignment==WD_ALIGN_PARAGRAPH.LEFT for p in commands)
    second=next(p for p in doc.paragraphs if p.text=='Experiment No. 2')
    assert second.paragraph_format.page_break_before is True
    figures=[p for p in doc.paragraphs if 'Fig ' in p.text]
    assert figures and all(p.paragraph_format.keep_together for p in figures)
